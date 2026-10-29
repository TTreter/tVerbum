"""
Módulo de Extração e Análise de Pautas / Ordem do Dia (tVerbum v2.0+).
Extrai proposições, autores e ementas a partir de arquivos PDF, Word (.docx) ou texto livre
para ancoragem obrigatória do modelo de IA, eliminando alucinações em números de leis.
"""

import re
import io
import logging
from typing import List, Dict, Any, Optional
from pathlib import Path
from pypdf import PdfReader
from docx import Document

from .models import VotedMatter

logger = logging.getLogger("tverbum.agenda")

# Padrões de proposições legislativas usuais no Brasil
PROPOSITION_PATTERNS = [
    r"(Projeto de Lei Complementar\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Projeto de Lei Ordin[aá]ria\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Projeto de Lei\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Projeto de Decreto Legislativo\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Projeto de Resolu[cç][aã]o\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Proposta de Emenda [aà] Lei Org[aâ]nica\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Proposta de Emenda [aà] Constitui[cç][aã]o\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Requerimento\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Indica[cç][aã]o\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Mo[cç][aã]o\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Veto\s*(?:Total|Parcial)?\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Of[ií]cio\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Pedido de Informa[cç][aã]o\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
    r"(Pedido de Provid[eê]ncia\s*(?:n[º°.]?\s*)?\d+[/-]\d{4})",
]

COMPILED_PROPOSITION_RE = re.compile(
    r"\b(?:" + "|".join(PROPOSITION_PATTERNS) + r")\b",
    re.IGNORECASE,
)


class AgendaParser:
    """Extrai proposições legislativas e contexto temático de pautas e ordens do dia."""

    @classmethod
    def extract_text_from_pdf(cls, file_bytes: bytes) -> str:
        """Lê o texto integral de um arquivo PDF via pypdf."""
        try:
            reader = PdfReader(io.BytesIO(file_bytes))
            pages_text = []
            for idx, page in enumerate(reader.pages):
                text = page.extract_text() or ""
                if text.strip():
                    pages_text.append(text.strip())
            return "\n\n".join(pages_text)
        except Exception as e:
            logger.error(f"Erro ao extrair texto de PDF da pauta: {e}")
            raise ValueError(f"Não foi possível ler o arquivo PDF da pauta: {str(e)}")

    @classmethod
    def extract_text_from_docx(cls, file_bytes: bytes) -> str:
        """Lê o texto integral de um documento Word (.docx) via python-docx."""
        try:
            doc = Document(io.BytesIO(file_bytes))
            paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
            for table in doc.tables:
                for row in table.rows:
                    row_text = " | ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                    if row_text:
                        paragraphs.append(row_text)
            return "\n".join(paragraphs)
        except Exception as e:
            logger.error(f"Erro ao extrair texto de DOCX da pauta: {e}")
            raise ValueError(f"Não foi possível ler o arquivo Word (.docx) da pauta: {str(e)}")

    @classmethod
    def parse_propositions_from_text(cls, text: str) -> List[Dict[str, str]]:
        """
        Analisa o texto da pauta e extrai deterministicamente uma lista estruturada de proposições,
        tentando identificar o autor e a ementa associada a cada uma.
        """
        if not text:
            return []

        # Normaliza espaços
        lines = [line.strip() for line in text.splitlines() if line.strip()]
        cleaned_text = "\n".join(lines)

        # Encontra todas as ocorrências de proposições com suas posições
        matches = list(COMPILED_PROPOSITION_RE.finditer(cleaned_text))
        if not matches:
            return []

        extracted_matters: List[Dict[str, str]] = []

        for i, m in enumerate(matches):
            raw_matter = m.group(0).strip()
            # Padroniza a nomenclatura (ex: "Projeto de Lei nº 48/2026")
            clean_matter = re.sub(r"\s+", " ", raw_matter)
            clean_matter = re.sub(r"\bno\b|\bnº\b|\bn°\b", "nº", clean_matter, flags=re.IGNORECASE)
            if "nº" not in clean_matter and "/" in clean_matter:
                clean_matter = re.sub(r"(\d+/\d{4})", r"nº \1", clean_matter)

            # O bloco de texto desta proposição vai até o início da próxima (ou até 600 caracteres)
            start_pos = m.end()
            end_pos = matches[i + 1].start() if i + 1 < len(matches) else min(len(cleaned_text), start_pos + 600)
            block = cleaned_text[start_pos:end_pos].strip()

            # Tenta capturar autoria
            author = ""
            author_match = re.search(
                r"(?:Autoria|Autor|De autoria|Iniciativa|Proponente):\s*([^\n;.]+)",
                block,
                re.IGNORECASE,
            )
            if author_match:
                author = author_match.group(1).strip()
            elif re.search(r"Poder Executivo", block, re.IGNORECASE):
                author = "Poder Executivo"

            # Tenta capturar ementa
            summary = ""
            summary_match = re.search(
                r"(?:Ementa|Assunto|S[uú]mula):\s*([^\n]+(?:\n[^\n]+)?)",
                block,
                re.IGNORECASE,
            )
            if summary_match:
                summary = re.sub(r"\s+", " ", summary_match.group(1).strip())
            else:
                # Pega as primeiras linhas do bloco após a proposição como ementa descritiva
                first_lines = [l for l in block.splitlines() if l and not l.lower().startswith("aut")][:2]
                summary = " ".join(first_lines)[:250].strip()

            extracted_matters.append({
                "matter": clean_matter,
                "author": author,
                "summary": summary,
            })

        return extracted_matters

    @classmethod
    def format_agenda_context(cls, matters: List[Dict[str, str]]) -> str:
        """Formata a lista de proposições em texto claro para ser injetado no prompt da IA."""
        lines = []
        for idx, m in enumerate(matters, start=1):
            line = f"{idx}. {m['matter']}"
            if m.get("author"):
                line += f" (Autoria: {m['author']})"
            if m.get("summary"):
                line += f" - Ementa: {m['summary']}"
            lines.append(line)
        return "\n".join(lines)
