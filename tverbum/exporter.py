"""
Exportador de atas geradas para diferentes formatos de arquivo (PDF, DOCX, Markdown, TXT).
Inclui diagramação oficial com suporte a Brasão/Logotipo, cabeçalho institucional e texto justificado.
"""

import os
import io
import base64
import urllib.request
from typing import Optional
from docx import Document
from docx.shared import Pt, Inches
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT

from reportlab.lib.pagesizes import A4
from reportlab.lib.units import cm
from reportlab.lib.enums import TA_JUSTIFY, TA_CENTER, TA_LEFT
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    Table,
    TableStyle,
    Image as RLImage,
)
from reportlab.lib import colors

from .models import MinuteOutput


class MinuteExporter:
    """Exporta o objeto MinuteOutput para arquivos físicos (.pdf, .docx, .md, .txt)."""

    @staticmethod
    def _resolve_logo_bytes(logo_url: Optional[str]) -> Optional[bytes]:
        """Resolve imagem de logotipo a partir de Data-URI base64, caminho de arquivo local ou URL HTTP."""
        if not logo_url:
            return None
        try:
            logo_url = logo_url.strip()
            if logo_url.startswith("data:image"):
                header, b64data = logo_url.split(",", 1)
                return base64.b64decode(b64data)
            if os.path.exists(logo_url):
                with open(logo_url, "rb") as f:
                    return f.read()
            if logo_url.startswith("http://") or logo_url.startswith("https://"):
                with urllib.request.urlopen(logo_url, timeout=5) as resp:
                    return resp.read()
        except Exception:
            return None
        return None

    @staticmethod
    def to_txt(output: MinuteOutput, file_path: str) -> str:
        """Salva a ata em formato de texto simples (.txt)."""
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(output.full_text)
        return file_path

    @staticmethod
    def to_markdown(output: MinuteOutput, file_path: str) -> str:
        """Salva a ata em formato Markdown (.md)."""
        md_lines = []
        if output.header:
            md_lines.append("```")
            md_lines.append(output.header)
            md_lines.append("```\n")
        else:
            md_lines.append(f"# {output.title}\n")

        md_lines.append(output.body)

        if output.signatures:
            md_lines.append("\n\n```")
            md_lines.append(output.signatures.strip())
            md_lines.append("```")

        content = "\n".join(md_lines)
        with open(file_path, "w", encoding="utf-8") as f:
            f.write(content)
        return file_path

    @classmethod
    def to_docx(cls, output: MinuteOutput, file_path: str) -> str:
        """
        Salva a ata em formato Microsoft Word (.docx) com diagramação oficial:
        - Suporte a Brasão/Logotipo no cabeçalho
        - Alinhamento justificado no corpo do texto
        - Suporte a parágrafo único ou múltiplos parágrafos
        - Tabela de assinaturas alinhada
        """
        doc = Document()

        for section in doc.sections:
            section.top_margin = Inches(0.9)
            section.bottom_margin = Inches(0.9)
            section.left_margin = Inches(1.1)
            section.right_margin = Inches(1.1)

        logo_bytes = cls._resolve_logo_bytes(output.logo_url)

        # Cabeçalho Institucional
        if logo_bytes and (output.state or output.organization_name):
            hdr_table = doc.add_table(rows=1, cols=2)
            hdr_table.alignment = WD_TABLE_ALIGNMENT.CENTER
            cell_logo = hdr_table.rows[0].cells[0]
            cell_logo.width = Inches(1.2)
            p_logo = cell_logo.paragraphs[0]
            p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run_logo = p_logo.add_run()
            run_logo.add_picture(io.BytesIO(logo_bytes), width=Inches(0.95))

            cell_text = hdr_table.rows[0].cells[1]
            cell_text.width = Inches(5.1)
            p_text = cell_text.paragraphs[0]
            p_text.alignment = WD_ALIGN_PARAGRAPH.LEFT
            if output.state:
                r_st = p_text.add_run(output.state.upper() + "\n")
                r_st.font.name = "Arial"
                r_st.font.size = Pt(12)
            if output.organization_name:
                r_org = p_text.add_run(output.organization_name.upper() + "\n")
                r_org.font.name = "Arial"
                r_org.font.size = Pt(13)
                r_org.font.bold = True
            if output.organization_details:
                r_det = p_text.add_run(output.organization_details.strip())
                r_det.font.name = "Arial"
                r_det.font.size = Pt(8.5)

            doc.add_paragraph()
            p_title = doc.add_paragraph()
            p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
            r_title = p_title.add_run(output.title.upper())
            r_title.font.name = "Arial"
            r_title.font.size = Pt(12)
            r_title.font.bold = True
            p_title.paragraph_format.space_after = Pt(12)

        elif output.header:
            for line in output.header.splitlines():
                line_clean = line.strip()
                if not line_clean:
                    continue
                p = doc.add_paragraph()
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                run = p.add_run(line_clean)
                run.font.name = "Arial"
                is_title_or_org = line_clean.startswith("ATA") or line_clean == output.organization_name.upper()
                run.font.size = Pt(12 if is_title_or_org else 10)
                run.font.bold = is_title_or_org
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.space_before = Pt(0)
            doc.add_paragraph()
        else:
            p = doc.add_paragraph()
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            run = p.add_run(output.title)
            run.font.name = "Arial"
            run.font.size = Pt(12)
            run.font.bold = True
            p.paragraph_format.space_after = Pt(12)

        # Corpo da Ata (suporta 1 ou N parágrafos)
        paragraphs = [p.strip() for p in output.body.split("\n\n") if p.strip()]
        if not paragraphs:
            paragraphs = [output.body.strip()]

        for idx, para_text in enumerate(paragraphs):
            body_p = doc.add_paragraph()
            body_p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
            body_run = body_p.add_run(para_text)
            body_run.font.name = "Arial"
            body_run.font.size = Pt(11)
            body_p.paragraph_format.line_spacing = 1.35
            body_p.paragraph_format.space_after = Pt(10 if idx < len(paragraphs) - 1 else 18)

        # Assinaturas
        if output.signatures:
            doc.add_paragraph()
            table = doc.add_table(rows=2, cols=2)
            table.alignment = WD_TABLE_ALIGNMENT.CENTER

            row_lines = table.rows[0].cells
            row_lines[0].text = "____________________________________"
            row_lines[1].text = "____________________________________"
            row_lines[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row_lines[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

            row_labels = table.rows[1].cells
            row_labels[0].text = "Secretário"
            row_labels[1].text = "Presidente"
            row_labels[0].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER
            row_labels[1].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.CENTER

            for row in table.rows:
                for cell in row.cells:
                    for paragraph in cell.paragraphs:
                        for run in paragraph.runs:
                            run.font.name = "Arial"
                            run.font.size = Pt(10)

        doc.save(file_path)
        return file_path

    @classmethod
    def to_pdf(cls, output: MinuteOutput, file_path: str) -> str:
        """
        Salva a ata em formato PDF oficial diagramado:
        - Cabeçalho institucional com Brasão (se fornecido), Estado, Câmara e Contatos
        - Título centralizado em negrito
        - Corpo justificado com espaçamento formal
        - Linhas de assinatura para Secretário e Presidente
        """
        doc = SimpleDocTemplate(
            file_path,
            pagesize=A4,
            leftMargin=2.5 * cm,
            rightMargin=2.5 * cm,
            topMargin=2.0 * cm,
            bottomMargin=2.0 * cm,
            title=output.title,
        )

        styles = getSampleStyleSheet()
        style_state = ParagraphStyle(
            "StateHeader",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=13,
            leading=16,
            alignment=TA_LEFT,
            textColor=colors.black,
        )
        style_org = ParagraphStyle(
            "OrgHeader",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=15,
            leading=18,
            alignment=TA_LEFT,
            textColor=colors.black,
        )
        style_details = ParagraphStyle(
            "DetailsHeader",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=8.5,
            leading=11,
            alignment=TA_LEFT,
            textColor=colors.black,
        )
        style_center_hdr = ParagraphStyle(
            "CenterHeader",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=12,
            leading=15,
            alignment=TA_CENTER,
        )
        style_center_sub = ParagraphStyle(
            "CenterSub",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=9,
            leading=12,
            alignment=TA_CENTER,
        )
        style_title = ParagraphStyle(
            "MinuteTitle",
            parent=styles["Normal"],
            fontName="Helvetica-Bold",
            fontSize=12.5,
            leading=16,
            alignment=TA_CENTER,
            spaceBefore=12,
            spaceAfter=14,
        )
        style_body = ParagraphStyle(
            "MinuteBody",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=11,
            leading=16.5,
            alignment=TA_JUSTIFY,
            spaceAfter=10,
        )
        style_sig = ParagraphStyle(
            "MinuteSig",
            parent=styles["Normal"],
            fontName="Helvetica",
            fontSize=10.5,
            leading=14,
            alignment=TA_CENTER,
        )

        story = []
        logo_bytes = cls._resolve_logo_bytes(output.logo_url)

        # Construção do Cabeçalho
        if logo_bytes and (output.state or output.organization_name):
            img = RLImage(io.BytesIO(logo_bytes), width=2.3 * cm, height=2.3 * cm)
            right_flow = []
            if output.state:
                right_flow.append(Paragraph(output.state.upper(), style_state))
            if output.organization_name:
                right_flow.append(Paragraph(output.organization_name.upper(), style_org))
            if output.organization_details:
                for d_line in output.organization_details.splitlines():
                    if d_line.strip():
                        right_flow.append(Paragraph(d_line.strip(), style_details))

            hdr_table = Table([[img, right_flow]], colWidths=[2.8 * cm, 13.2 * cm])
            hdr_table.setStyle(
                TableStyle(
                    [
                        ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                        ("ALIGN", (0, 0), (0, 0), "CENTER"),
                        ("LEFTPADDING", (0, 0), (-1, -1), 0),
                        ("RIGHTPADDING", (0, 0), (-1, -1), 0),
                    ]
                )
            )
            story.append(hdr_table)
            story.append(Spacer(1, 0.4 * cm))
        else:
            if output.state:
                story.append(Paragraph(output.state.upper(), style_center_hdr))
            if output.organization_name:
                story.append(Paragraph(output.organization_name.upper(), style_center_hdr))
            if output.organization_details:
                for d_line in output.organization_details.splitlines():
                    if d_line.strip():
                        story.append(Paragraph(d_line.strip(), style_center_sub))
            if output.state or output.organization_name:
                story.append(Spacer(1, 0.3 * cm))

        # Título da Ata
        story.append(Paragraph(output.title.upper(), style_title))

        # Corpo da Ata
        paragraphs = [p.strip() for p in output.body.split("\n\n") if p.strip()]
        if not paragraphs:
            paragraphs = [output.body.strip()]

        for para_text in paragraphs:
            safe_text = (
                para_text.replace("&", "&amp;")
                .replace("<", "&lt;")
                .replace(">", "&gt;")
            )
            story.append(Paragraph(safe_text, style_body))

        # Bloco de Assinaturas
        if output.signatures:
            story.append(Spacer(1, 1.4 * cm))
            sig_data = [
                [
                    Paragraph("______________________________", style_sig),
                    Paragraph("______________________________", style_sig),
                ],
                [
                    Paragraph("Secretário", style_sig),
                    Paragraph("Presidente", style_sig),
                ],
            ]
            sig_table = Table(sig_data, colWidths=[8.0 * cm, 8.0 * cm])
            sig_table.setStyle(
                TableStyle(
                    [
                        ("ALIGN", (0, 0), (-1, -1), "CENTER"),
                        ("VALIGN", (0, 0), (-1, -1), "TOP"),
                        ("TOPPADDING", (0, 0), (-1, -1), 2),
                        ("BOTTOMPADDING", (0, 0), (-1, -1), 2),
                    ]
                )
            )
            story.append(sig_table)

        doc.build(story)
        return file_path
